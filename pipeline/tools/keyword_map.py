"""Keyword map: one primary keyword per indexable page, plus the Search Console queries (last 12
months) assigned to the page that should rank for them. Output: docs/keyword-map.xlsx (+ .csv)

Usage: python3 -I tools/keyword_map.py <site-content-dir> <out-dir>"""
import csv, glob, json, os, re, sys, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

CONTENT, OUT = sys.argv[1], sys.argv[2]
GSC = "/root/.claude/uploads/edc57836-ba21-5b6f-a1ee-e86b56115075/fc9d64e1-funplaysystems.com-Performance-on-Search-2026-10-10.xlsx"
HOST = "https://funplaysystems.com"
FACTS = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "briefs", "cities.json")))
EXCLUDE = {"/search-result.html", "/thank-you.html", "/thankyou.html", "/thank-you-lp.html", "/thank-you-campaign.html", "/404.html",
           "/generic.html", "/ads.html", "/campaign.html", "/privacy-policy.html", "/terms-of-service.html"}
CORE = {
    "/": ("playground equipment manufacturer in India", "outdoor playground equipment; children play equipment; FunPlay Systems; funplay"),
    "/Products.html": ("playground equipment products", "outdoor play equipment; outdoor gym equipment; garden furniture"),
    "/Projects.html": ("playground installation projects India", "school playground projects; society playground projects"),
    "/about-us.html": ("FunPlay Systems", "playground equipment company Vasai; playground makers since 2002"),
    "/blog.html": ("playground design and safety blog", "playground buying guide"),
    "/builders.html": ("playground equipment for builders", "society playground equipment; developer playground design"),
    "/school.html": ("school playground equipment", "school play equipment; preschool play equipment; playground for schools"),
    "/careers.html": ("FunPlay Systems careers", "jobs in Vasai"),
    "/clientele.html": ("FunPlay Systems clients", "playground equipment for schools builders government"),
    "/contact-us.html": ("playground equipment quote", "FunPlay Systems contact; FunPlay Vasai address"),
    "/corecraft.html": ("CoreCraft playground equipment", "multi play station; modular playground equipment"),
    "/faq.html": ("playground equipment FAQs", "playground safety questions"),
    "/funplay-originals.html": ("designer playground equipment", "FunPlay Originals; Red Dot playground"),
    "/gallery.html": ("playground installation photos", "playground gallery"),
    "/industry.html": ("playground equipment for schools, builders and parks", "commercial playground equipment"),
    "/nexlev-series.html": ("wooden play structures", "NexLev series; wooden playground equipment"),
    "/open-air-gym-equipment-manufacturer-and-supplier-in-india.html": ("open air gym equipment manufacturer in India", "outdoor gym equipment India; open gym equipment"),
    "/our-network.html": ("playground equipment supplier across India", "playground equipment near me"),
    "/outdoor-play-equipment.html": ("outdoor play equipment for kids", "outdoor playground equipment"),
    "/outdoor-playground-swing-sets-manufacturers-in-india.html": ("outdoor swing sets manufacturer in India", "playground swings; swing set for school"),
}
STOP = {"in", "the", "for", "and", "of", "a", "to", "near", "me", "best", "top", "india", "manufacturer", "manufacturers", "supplier", "suppliers",
        "equipment", "equipments", "playground", "outdoor", "play", "gym", "kids", "children", "company", "companies"}


def load_pages():
    recs = []
    for f in glob.glob(os.path.join(CONTENT, "pages", "*.json")):
        recs += json.load(open(f, encoding="utf-8"))
    return [r for r in recs if r["type"] in ("home", "page", "blog", "project", "series", "product", "city-gym", "city-playground")
            and r["path"] not in EXCLUDE and "noindex" not in ((r.get("seo") or {}).get("robots") or "")]


def place_of(r):
    f = FACTS.get(r["path"]) or {}
    p = f.get("city") or f.get("state")
    if not p:
        m = re.search(r"\bin ([^|<]+)$", r.get("h1") or "")
        p = m.group(1).strip() if m else None
    return p


def primary(r):
    p, t = r["path"], r["type"]
    if p in CORE:
        return CORE[p]
    name = (r.get("h1") or r["seo"].get("title", "")).strip()
    if t == "city-playground":
        pl = place_of(r)
        return (f"playground equipment manufacturer in {pl}", f"outdoor playground equipment {pl}; children play equipment {pl}; school playground equipment {pl}")
    if t == "city-gym":
        pl = place_of(r)
        return (f"outdoor gym equipment supplier in {pl}", f"outdoor gym equipment {pl}; open gym equipment {pl}; outdoor fitness equipment {pl}")
    if t == "product":
        fit = "fitness" in p or "gym" in p
        return (f"{name.lower()} {'outdoor gym equipment' if fit else 'playground equipment'}", f"{name.lower()}; FunPlay {name}")
    if t == "series":
        return (f"{name.lower()}", "playground equipment range")
    if t == "project":
        return (f"{(r.get('data') or {}).get('project', {}).get('title') or name} playground project".lower(), "playground installation")
    if t == "blog":
        return (re.sub(r"\s*[|:].*$", "", name).lower(), "")
    return (name.lower(), "")


def tokens(s):
    return [w for w in re.findall(r"[a-z0-9]+", s.lower()) if w not in STOP and len(w) > 2]


def main():
    pages = load_pages()
    wb = openpyxl.load_workbook(GSC, read_only=True)
    gp = collections.defaultdict(lambda: [0, 0, 0.0])
    for row in list(wb["Pages"].iter_rows(values_only=True))[1:]:
        if not row[0]:
            continue
        u = row[0].replace("https://www.funplaysystems.com", HOST).replace(HOST, "") or "/"
        a = gp[u]
        a[2] = (a[2] * a[1] + float(row[4] or 0) * float(row[2] or 0)) / max(1, a[1] + float(row[2] or 0))
        a[0] += int(row[1] or 0); a[1] += int(row[2] or 0)
    queries = [(r[0], int(r[1] or 0), int(r[2] or 0), float(r[4] or 0)) for r in list(wb["Queries"].iter_rows(values_only=True))[1:] if r[0]]

    rows, by_path = [], {}
    for r in pages:
        kw, sec = primary(r)
        c, i, pos = gp.get(r["path"], (0, 0, 0))
        d = {"URL": HOST + r["path"], "path": r["path"], "Page type": r["type"], "Primary keyword": kw, "Secondary keywords": sec,
             "Title tag": r["seo"].get("title", ""), "H1": (r.get("h1") or "").strip(), "GSC clicks (12m)": c, "GSC impressions (12m)": i,
             "GSC avg position": round(pos, 1) if i else "", "Matched queries": [], "Place": place_of(r) if r["type"].startswith("city") else ""}
        rows.append(d); by_path[r["path"]] = d

    # assign queries: brand -> home; city names -> that city page (gym vs playground by wording); product names -> product page
    places = collections.defaultdict(list)
    for d in rows:
        if d["Place"]:
            places[d["Place"].lower()].append(d)
    names = {d["H1"].lower(): d for d in rows if d["Page type"] == "product" and len(d["H1"]) > 3}
    for d in rows:
        if d["Page type"] == "product":
            names.setdefault(d["path"].rsplit("/", 1)[-1][:-5].replace("-", " "), d)
    assign = []
    for q, c, i, pos in queries:
        ql = q.lower()
        target, why = None, ""
        if re.search(r"\bfun ?play\b", ql) and not re.search(r"equipment|playground|gym|swing|slide|\bin\b", ql):
            target, why = by_path["/"], "brand"
        if not target:
            for pl, ds in sorted(places.items(), key=lambda kv: -len(kv[0])):
                if pl in ("india",):
                    continue  # national terms belong to the home page and the open-air gym page
                if pl and re.search(r"\b" + re.escape(pl) + r"\b", ql):
                    gym = any(w in ql for w in ("gym", "fitness", "exercise"))
                    pick = [x for x in ds if (x["Page type"] == "city-gym") == gym] or ds
                    target, why = pick[0], "place name"
                    break
        if not target:
            for n, d in sorted(names.items(), key=lambda kv: -len(kv[0])):
                if len(n) <= 4:
                    continue  # very short product names ("Fun") match too much
                if re.search(r"\b" + re.escape(n) + r"\b", ql):
                    target, why = d, "product name"
                    break
        if not target and len(ql.split()) == 1:
            import difflib
            hit = difflib.get_close_matches(ql, list(names), n=1, cutoff=0.75)
            if hit:
                target, why = names[hit[0]], "product name (misspelt)"
        if not target:
            qt = set(tokens(ql))
            best, score = None, 0
            for d in rows:
                if d["Page type"] in ("product", "project", "city-gym", "city-playground"):
                    continue  # products and places are matched by name above
                kt = set(tokens(d["Primary keyword"] + " " + d["Secondary keywords"]))
                s = len(qt & kt) / (len(qt) or 1) + (0.01 if d["Page type"] in ("home", "page") else 0)
                if s > score or (s == score and best and d["GSC impressions (12m)"] > best["GSC impressions (12m)"]):
                    best, score = d, s
            if best and score >= 0.5:
                target, why = best, "keyword overlap"
        if not target:
            gym = any(w in ql for w in ("gym", "fitness", "exercise"))
            target = by_path["/open-air-gym-equipment-manufacturer-and-supplier-in-india.html"] if gym else by_path["/"]
            why = "generic (category page)"
        target["Matched queries"].append((q, c, i))
        assign.append({"Query": q, "Clicks": c, "Impressions": i, "Avg position": round(pos, 1), "Assigned page": target["URL"], "Rule": why})

    os.makedirs(OUT, exist_ok=True)
    out = openpyxl.Workbook()
    ws = out.active; ws.title = "Keyword map"
    cols = ["URL", "Page type", "Primary keyword", "Secondary keywords", "Title tag", "H1", "GSC clicks (12m)", "GSC impressions (12m)", "GSC avg position", "Top GSC queries for this page"]
    ws.append(cols)
    order = {"home": 0, "page": 1, "series": 2, "product": 3, "blog": 4, "project": 5, "city-playground": 6, "city-gym": 7}
    rows.sort(key=lambda d: (order.get(d["Page type"], 9), -d["GSC impressions (12m)"], d["URL"]))
    for d in rows:
        top = "; ".join(q for q, _, _ in sorted(d["Matched queries"], key=lambda x: (-x[1], -x[2]))[:5])
        ws.append([d[c] for c in cols[:-1]] + [top])
    ws2 = out.create_sheet("GSC query assignment")
    acols = ["Query", "Clicks", "Impressions", "Avg position", "Assigned page", "Rule"]
    ws2.append(acols)
    for a in assign:
        ws2.append([a[c] for c in acols])
    ws3 = out.create_sheet("Notes")
    for line in [
        "How to read this workbook",
        "Keyword map: one primary keyword per indexable page (title, H1 and copy are written around it). Secondary keywords are supporting terms.",
        "GSC figures are the 12-month Search Console export dated 10 Oct 2026, with www and non-www URLs combined.",
        "GSC query assignment: each of the top 1,000 queries is assigned to the page that should rank for it, by brand, place name, product name or keyword overlap.",
        "City pages follow one pattern: 'playground equipment manufacturer in <place>' and 'outdoor gym equipment supplier in <place>', so each city has one playground page and one gym page and they do not compete.",
        "Change any primary keyword here first, then update the page title, H1 and intro to match.",
    ]:
        ws3.append([line])
    for w in (ws, ws2):
        for cell in w[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="D74951")
        w.freeze_panes = "A2"
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col[:200]) + 2))
    ws3.column_dimensions["A"].width = 140
    ws3["A1"].font = Font(bold=True)
    out.save(os.path.join(OUT, "keyword-map.xlsx"))
    with open(os.path.join(OUT, "keyword-map.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(cols)
        for d in rows:
            w.writerow([d[c] for c in cols[:-1]] + ["; ".join(q for q, _, _ in sorted(d["Matched queries"], key=lambda x: (-x[1], -x[2]))[:5])])
    print("pages", len(rows), "queries", len(assign), collections.Counter(a["Rule"] for a in assign))


main()
